"""Create text, code, and Word files directly, without opening an editor or running the code."""

import ast
from pathlib import Path

OUTPUT_DIR = Path(__file__).resolve().parent / "output"


# Save a new file and read it back; existing files are never overwritten.
def write_file(path, content):
    if not isinstance(path, str) or not path.strip() or not isinstance(content, str):
        raise ValueError("Give a file path and the text to save.")
    target = Path(path).expanduser()
    if not target.is_absolute():
        target = OUTPUT_DIR / target
    target = target.resolve()
    # Word files use a document package; the other supported formats store plain text.
    if target.suffix.lower() not in ("", ".docx", ".py", ".txt", ".md", ".json", ".csv", ".html",
                                     ".css", ".js", ".ts", ".yaml", ".yml", ".xml", ".toml", ".ini", ".log"):
        raise ValueError("Choose a supported text, code, or Word file, such as .txt, .py, or .docx.")
    if target.suffix.lower() == ".py":
        ast.parse(content)  # Check Python syntax without executing anything.
    expected = content.replace("\r\n", "\n").replace("\r", "\n")
    if target.suffix.lower() == ".docx":
        from docx import Document

        document = Document()
        for line in expected.split("\n"):
            document.add_paragraph(line)  # Each line becomes a Word paragraph.
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("xb") as file:
            document.save(file)  # Create a real Word package, not renamed text.
    else:
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("x", encoding="utf-8", newline="") as file:
            file.write(content)  # x means create only; fail if the file already exists.
    if read_file(target) != expected:
        raise RuntimeError("The saved file could not be verified. Check it before trying again.")
    return str(target)


# Read saved text in the right format so both file creation and completion can verify it.
def read_file(path):
    path = Path(path)
    if path.suffix.lower() == ".docx":
        from docx import Document

        return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    return path.read_text(encoding="utf-8")
